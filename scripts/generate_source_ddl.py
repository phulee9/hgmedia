#!/usr/bin/env python3
"""Generate local source-replica DDL for the database tables used by the EL pipeline.

The generator intentionally creates only physical source columns observed in the
latest staging sample export. Data Dictionary metadata is used to determine
source table identity and SQL types whenever the dictionary documents the field.
Fields not covered by the dictionary are inferred from sample values and are
recorded in the coverage report.

Usage:
  python scripts/generate_source_ddl.py \
      --dictionary /path/to/HG_Media_Source_Data_Dictionary_v2.xlsm
"""
from __future__ import annotations

import argparse
import csv
import json
import re
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_EXPORT_DIR = REPO_ROOT / "exports" / "hgmediadb_samples_1000_20260828_000045"
STAGING_META_COLUMNS = {
    "_source_id", "_source_connection", "_batch_id", "_loaded_at", "_wm", "_year",
}

SHEET_BY_CONNECTION = {
    "odoo_pg": "Odoo",
    "hg_stock": "HG_Stock",
    "editing_management": "Editing",
    "record_survey": "Cham_nhac",
    "channel_channel": "QLK_Channel",
    "channel_network": "QLK_Network",
    "channel_organization": "QLK_Organization",
    "channel_project": "QLK_Project",
    "channel_relationship": "QLK_Relationship",
}

SQLSERVER_DATABASE_BY_CONNECTION = {
    "editing_management": "editing-management",
    "record_survey": "record-survey",
    "channel_channel": "channel.service.channel",
    "channel_network": "channel.service.network",
    "channel_organization": "channel.service.organization",
    "channel_project": "channel.service.project",
    "channel_relationship": "channel.service.relationship",
}

POSTGRES_DATABASE_BY_CONNECTION = {"odoo_pg": "odoo", "hg_stock": "hgstock"}


@dataclass(frozen=True)
class DictField:
    sheet: str
    table: str
    column: str
    data_type: str
    description: str
    nullable: str
    data_length: str


def norm(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def table_aliases(connection: str, table: str) -> list[str]:
    aliases = [table]
    if connection == "odoo_pg":
        aliases.append(table.replace("_", "."))
    elif connection == "editing_management":
        aliases.extend({
            "Editings": ["Editing"],
            "Resources": ["Resource"],
            "Resource_Editings": ["Resource_Editing"],
            "User": ["User"],
        }.get(table, []))
    elif connection == "channel_network" and table == "Network":
        aliases.append("Networks")
    elif connection == "channel_relationship" and table == "channel_department":
        aliases.extend(["Channel_Department", "ChannelDepartment"])
    return aliases


def load_dictionary(path: Path) -> dict[tuple[str, str, str], DictField]:
    wb = load_workbook(path, read_only=False, data_only=True, keep_vba=False)
    result: dict[tuple[str, str, str], DictField] = {}
    for sheet in wb.sheetnames:
        if sheet == "00_Muc_luc":
            continue
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) < 12:
                continue
            table, column = row[6], row[7]
            if table is None or column is None:
                continue
            field = DictField(
                sheet=sheet,
                table=str(table).strip(),
                column=str(column).strip(),
                data_type="" if row[8] is None else str(row[8]).strip(),
                description="" if row[9] is None else str(row[9]).strip(),
                nullable="" if row[10] is None else str(row[10]).strip(),
                data_length="" if row[11] is None else str(row[11]).strip(),
            )
            result[(sheet, norm(field.table), norm(field.column))] = field
    return result


def find_dict_field(index, connection: str, table: str, column: str) -> DictField | None:
    sheet = SHEET_BY_CONNECTION[connection]
    for alias in table_aliases(connection, table):
        hit = index.get((sheet, norm(alias), norm(column)))
        if hit:
            return hit
    return None


def sample_values(csv_path: Path, column: str, limit: int = 1000) -> list[str]:
    df = pd.read_csv(csv_path, dtype=str, keep_default_na=False, nrows=limit, low_memory=False)
    if column not in df.columns:
        return []
    vals = []
    for value in df[column].tolist():
        value = str(value).strip()
        if value and value.lower() not in {"nan", "none", "null", "nat"}:
            vals.append(value)
    return vals


def looks_datetime(values: list[str], column: str) -> bool:
    compact = norm(column)
    if compact.endswith("id") or compact.endswith("ids"):
        return False
    date_hint = (
        any(token in compact for token in (
            "createddate", "updateddate", "deleteddate", "publisheddate",
            "reportingdate", "effectivedate", "expiredate", "expirationdate",
            "birthdate", "birthday",
        ))
        or compact.endswith(("date", "time", "datetime", "utc", "timestamp"))
        or compact.startswith(("date", "time"))
    )
    if not date_hint:
        return False
    if not values:
        return True
    parsed = pd.to_datetime(pd.Series(values[:100]), errors="coerce", utc=True)
    return float(parsed.notna().mean()) >= 0.85


def infer_generic(values: list[str], column: str) -> str:
    compact = norm(column)
    lowered = {v.lower() for v in values[:500]}
    if any(token in compact for token in (
        "phone", "email", "username", "name", "title", "code", "url",
        "path", "hash", "barcode", "postal", "zipcode", "address",
    )):
        return "text"
    if looks_datetime(values, column):
        return "datetime"
    if compact.startswith(("is", "has", "can", "enable", "active")) and values and lowered <= {
        "true", "false", "t", "f", "1", "0", "yes", "no", "y", "n"
    }:
        return "bool"
    if values and lowered <= {"true", "false", "t", "f"}:
        return "bool"
    int_re = re.compile(r"^-?\d+$")
    dec_re = re.compile(r"^-?(?:\d+\.\d+|\d+\.?\d*[eE][+-]?\d+)$")
    if values and all(int_re.match(v) for v in values[:500]):
        return "bigint"
    if values and all(int_re.match(v) or dec_re.match(v) for v in values[:500]):
        return "decimal"
    return "text"


def dict_category(raw: str) -> str | None:
    s = raw.strip().lower()
    if not s:
        return None
    if "datetimeoffset" in s:
        return "datetimeoffset"
    if "datetime" in s:
        return "datetime"
    if s == "date" or s.startswith("date "):
        return "date"
    if "guid" in s or "uniqueidentifier" in s:
        return "guid"
    if "bool" in s or s == "bit" or "boolean" in s:
        return "bool"
    if "decimal" in s or "numeric" in s:
        return "decimal"
    if re.search(r"(^|\W)(float|double)(\W|$)", s):
        return "float"
    if any(token in s for token in ("bigint", "ulong", "long")):
        return "bigint"
    if re.search(r"(^|\W)int(\W|$)", s):
        return "bigint"
    if "binary" in s or "bytea" in s:
        return "binary"
    if s.startswith("m2o"):
        return "bigint"
    if s.startswith("m2m") or s.startswith("o2m") or "list~" in s:
        return "text"
    if any(token in s for token in (
        "string", "nvarchar", "varchar", "text", "html", "selection", "char",
        "resourcetype", "editingsoftware", "enum",
    )):
        return "text"
    return None


def sql_type(dialect: str, dict_type: str, values: list[str], column: str) -> tuple[str, str]:
    category = dict_category(dict_type) if dict_type else None
    provenance = "dictionary" if category else "sample_inferred"
    if category is None:
        category = infer_generic(values, column)

    if dialect == "postgresql":
        return {
            "datetimeoffset": "TIMESTAMPTZ", "datetime": "TIMESTAMP", "date": "DATE",
            "guid": "TEXT", "bool": "BOOLEAN", "decimal": "NUMERIC(38,10)",
            "float": "DOUBLE PRECISION", "bigint": "BIGINT", "binary": "BYTEA",
            "text": "TEXT",
        }.get(category, "TEXT"), provenance

    length_match = re.search(r"(?:string|nvarchar|varchar)\s*\(\s*(\d+)\s*\)", dict_type or "", flags=re.I)
    text_type = f"NVARCHAR({length_match.group(1)})" if length_match else "NVARCHAR(MAX)"
    return {
        "datetimeoffset": "DATETIMEOFFSET", "datetime": "DATETIME2", "date": "DATE",
        "guid": "UNIQUEIDENTIFIER", "bool": "BIT", "decimal": "DECIMAL(38,10)",
        "float": "FLOAT", "bigint": "BIGINT", "binary": "VARBINARY(MAX)",
        "text": text_type,
    }.get(category, "NVARCHAR(MAX)"), provenance


def pg_quote(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def ms_quote(name: str) -> str:
    return "[" + name.replace("]", "]]" ) + "]"


def read_headers(csv_path: Path) -> list[str]:
    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        return next(csv.reader(handle))


def build_table_ddl(dialect: str, schema: str, table: str, columns: list[tuple[str, str]]) -> str:
    if dialect == "postgresql":
        q = pg_quote
        col_sql = ",\n".join(f"    {q(name)} {typ} NULL" for name, typ in columns)
        return (
            f"CREATE SCHEMA IF NOT EXISTS {q(schema)};\n"
            f"DROP TABLE IF EXISTS {q(schema)}.{q(table)} CASCADE;\n"
            f"CREATE TABLE {q(schema)}.{q(table)} (\n{col_sql}\n);\n"
        )
    q = ms_quote
    col_sql = ",\n".join(f"    {q(name)} {typ} NULL" for name, typ in columns)
    return (
        f"IF OBJECT_ID(N'{q(schema)}.{q(table)}', N'U') IS NOT NULL DROP TABLE {q(schema)}.{q(table)};\nGO\n"
        f"CREATE TABLE {q(schema)}.{q(table)} (\n{col_sql}\n);\nGO\n"
    )


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dictionary", required=True, type=Path)
    ap.add_argument("--config", type=Path, default=REPO_ROOT / "config" / "db_sources.yaml")
    ap.add_argument("--export-dir", type=Path, default=DEFAULT_EXPORT_DIR)
    ap.add_argument("--output", type=Path, default=REPO_ROOT / "source_replicas" / "generated")
    args = ap.parse_args()

    dictionary = load_dictionary(args.dictionary)
    sources = yaml.safe_load(args.config.read_text(encoding="utf-8"))["db_sources"]
    args.output.mkdir(parents=True, exist_ok=True)
    (args.output / "postgres").mkdir(exist_ok=True)
    (args.output / "sqlserver").mkdir(exist_ok=True)
    report_dir = args.output.parent / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)

    pg_chunks: dict[str, list[str]] = {"odoo_pg": [], "hg_stock": []}
    ms_chunks: dict[str, list[str]] = {k: [] for k in SQLSERVER_DATABASE_BY_CONNECTION}
    coverage_rows = []
    manifest = []

    for src in sources:
        connection = src["connection"]
        table = src["source_table"]
        schema = src.get("schema", "dbo")
        target = src["target_staging_table"]
        csv_path = args.export_dir / f"{target}.csv"
        if not csv_path.exists():
            raise FileNotFoundError(f"Missing staging sample for {src['source_id']}: {csv_path}")

        headers = [c for c in read_headers(csv_path) if c not in STAGING_META_COLUMNS]
        dialect = "postgresql" if connection in POSTGRES_DATABASE_BY_CONNECTION else "sqlserver"
        typed_columns = []
        dict_hits = 0
        for col in headers:
            meta = find_dict_field(dictionary, connection, table, col)
            values = sample_values(csv_path, col)
            typ, provenance = sql_type(dialect, meta.data_type if meta else "", values, col)
            typed_columns.append((col, typ))
            if meta and provenance == "dictionary":
                dict_hits += 1
            coverage_rows.append({
                "source_id": src["source_id"],
                "connection": connection,
                "dictionary_sheet": SHEET_BY_CONNECTION[connection],
                "source_table": table,
                "source_column": col,
                "dictionary_table": meta.table if meta else "",
                "dictionary_column": meta.column if meta else "",
                "dictionary_type": meta.data_type if meta else "",
                "generated_sql_type": typ,
                "type_provenance": provenance,
                "dictionary_description": meta.description if meta else "",
            })

        ddl = build_table_ddl(dialect, schema, table, typed_columns)
        header = (
            f"-- source_id: {src['source_id']}\n"
            f"-- dictionary sheet: {SHEET_BY_CONNECTION[connection]}\n"
            f"-- physical columns: {len(headers)}; dictionary-typed columns: {dict_hits}\n"
        )
        if dialect == "postgresql":
            pg_chunks[connection].append(header + ddl)
            database = POSTGRES_DATABASE_BY_CONNECTION[connection]
        else:
            ms_chunks[connection].append(header + ddl)
            database = SQLSERVER_DATABASE_BY_CONNECTION[connection]

        manifest.append({
            "source_id": src["source_id"],
            "connection": connection,
            "dialect": dialect,
            "database": database,
            "schema": schema,
            "table": table,
            "target_staging_table": target,
            "fixture_csv": str(csv_path.relative_to(REPO_ROOT)),
            "source_columns": headers,
            "column_types": {name: typ for name, typ in typed_columns},
            "dictionary_typed_columns": dict_hits,
        })

    for connection, chunks in pg_chunks.items():
        out = args.output / "postgres" / f"{connection}.sql"
        out.write_text("\n\n".join(chunks) + "\n", encoding="utf-8")

    for connection, chunks in ms_chunks.items():
        db = SQLSERVER_DATABASE_BY_CONNECTION[connection]
        prefix = (
            f"IF DB_ID(N'{db}') IS NULL CREATE DATABASE {ms_quote(db)};\nGO\n"
            f"USE {ms_quote(db)};\nGO\n"
            "IF SCHEMA_ID(N'dbo') IS NULL EXEC('CREATE SCHEMA [dbo]');\nGO\n"
        )
        out = args.output / "sqlserver" / f"{connection}.sql"
        out.write_text(prefix + "\n\n" + "\n\n".join(chunks) + "\n", encoding="utf-8")

    with (report_dir / "source_dictionary_coverage.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(coverage_rows[0].keys()))
        writer.writeheader()
        writer.writerows(coverage_rows)

    (args.output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    total_columns = len(coverage_rows)
    typed = sum(1 for r in coverage_rows if r["type_provenance"] == "dictionary")
    print(f"Generated DDL for {len(manifest)} source tables")
    print(f"Dictionary typed columns: {typed}/{total_columns} ({typed/total_columns:.1%})")
    print(f"Output: {args.output}")


if __name__ == "__main__":
    main()
